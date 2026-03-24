"""Step implementations for numbering-related features."""

import os
import tempfile

from behave import given, then, when

from docx import Document

from helpers import test_docx

# given ===================================================


@given("a document having a numbering part")
def given_a_document_having_a_numbering_part(context):
    context.document = Document(test_docx("num-having-numbering-part"))


@given("a new document")
def given_a_new_document(context):
    context.document = Document()


@given("a new document with a custom numbering definition")
def given_a_new_document_with_custom_numbering(context):
    doc = Document()
    numbering_part = doc.part.numbering_part

    abstract_num = numbering_part.add_abstract_num("singleLevel")
    abstract_num.add_lvl(
        ilvl=0,
        num_fmt="decimal",
        lvl_text="%1.",
        indent_left=720,
        indent_hanging=360,
    )
    num = numbering_part.add_num(abstract_num.abstractNumId)

    p = doc.add_paragraph("Numbered item")
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_numId().val = num.numId
    numPr.get_or_add_ilvl().val = 0

    context.document = doc
    context.num_id = num.numId


# when ====================================================


@when("I get the numbering part from the document")
def when_get_numbering_part_from_document(context):
    document = context.document
    context.numbering_part = document.part.numbering_part


@when("I create a decimal numbering definition")
def when_create_decimal_numbering(context):
    doc = context.document
    numbering_part = doc.part.numbering_part

    abstract_num = numbering_part.add_abstract_num("singleLevel")
    abstract_num.add_lvl(
        ilvl=0,
        num_fmt="decimal",
        lvl_text="%1.",
        indent_left=720,
        indent_hanging=360,
    )
    num = numbering_part.add_num(abstract_num.abstractNumId)
    context.num_id = num.numId


@when("I apply the numbering to a paragraph")
def when_apply_numbering_to_paragraph(context):
    doc = context.document
    p = doc.add_paragraph("Numbered paragraph")
    pPr = p._element.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_numId().val = context.num_id
    numPr.get_or_add_ilvl().val = 0
    context.numbered_paragraph = p


@when("I save and reload the document")
def when_save_and_reload(context):
    doc = context.document
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        doc.save(tmp_path)
        context.document = Document(tmp_path)
    finally:
        os.unlink(tmp_path)


# then =====================================================


@then("the numbering part has the expected numbering definitions")
def then_numbering_part_has_expected_numbering_definitions(context):
    numbering_part = context.numbering_part
    assert len(numbering_part.numbering_definitions) == 10


@then("the paragraph has numbering applied")
def then_paragraph_has_numbering(context):
    p = context.numbered_paragraph
    numPr = p._element.pPr.numPr
    assert numPr is not None
    assert numPr.numId is not None


@then("the custom numbering definition is preserved")
def then_custom_numbering_preserved(context):
    doc = context.document
    numbering_part = doc.part.numbering_part
    numbering_elm = numbering_part.numbering_elm

    # Should have at least 3 abstract nums (2 defaults + 1 custom)
    assert len(numbering_elm.abstractNum_lst) >= 3

    # Find our custom one (the last one added, with highest abstractNumId)
    custom_an = max(numbering_elm.abstractNum_lst, key=lambda an: an.abstractNumId)
    assert len(custom_an.lvl_lst) >= 1
    assert custom_an.lvl_lst[0].numFmt_val == "decimal"
    assert custom_an.lvl_lst[0].lvlText_val == "%1."
