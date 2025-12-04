"""Test structural equivalence for generated test documents.

These tests create documents with various feature combinations and verify
they maintain structural equivalence through a read-save-read cycle.
"""

from pathlib import Path

import pytest

from docx import Document

from tests.roundtrip.document_generators import DOCUMENT_GENERATORS
from tests.roundtrip.structural_equivalence import (
    compare_documents,
    test_roundtrip_equivalence,
)


class DescribeGeneratedDocumentRoundTrip:
    """Tests for round-trip structural equivalence of generated documents."""

    @pytest.mark.parametrize(
        "generator_name",
        list(DOCUMENT_GENERATORS.keys()),
        ids=list(DOCUMENT_GENERATORS.keys()),
    )
    def it_preserves_generated_document_structure(
        self, generator_name: str, tmp_path: Path
    ):
        """Test that generated documents preserve structure through round-trip."""
        # Create the document
        generator = DOCUMENT_GENERATORS[generator_name]
        doc = generator()

        # Save initial version
        initial_path = tmp_path / f"{generator_name}_initial.docx"
        doc.save(str(initial_path))

        # Test round-trip
        result = test_roundtrip_equivalence(initial_path, tmp_path)

        assert result.is_equivalent, (
            f"Generated document '{generator_name}' has semantic differences:\n"
            f"  File diffs: {result.file_diffs}\n"
            f"  Semantic diffs: {result.semantic_diffs}"
        )


class DescribeDoubleRoundTrip:
    """Test that documents survive multiple round-trips."""

    @pytest.mark.parametrize(
        "generator_name",
        list(DOCUMENT_GENERATORS.keys()),
        ids=list(DOCUMENT_GENERATORS.keys()),
    )
    def it_survives_double_roundtrip(self, generator_name: str, tmp_path: Path):
        """Test that documents are equivalent after two round-trips."""
        # Create the document
        generator = DOCUMENT_GENERATORS[generator_name]
        doc = generator()

        # Save initial version
        path1 = tmp_path / f"{generator_name}_v1.docx"
        doc.save(str(path1))

        # First round-trip
        doc2 = Document(str(path1))
        path2 = tmp_path / f"{generator_name}_v2.docx"
        doc2.save(str(path2))

        # Second round-trip
        doc3 = Document(str(path2))
        path3 = tmp_path / f"{generator_name}_v3.docx"
        doc3.save(str(path3))

        # Compare v2 to v3 (both are round-tripped versions)
        result = compare_documents(path2, path3)

        assert result.is_equivalent, (
            f"Document '{generator_name}' differs after double round-trip:\n"
            f"  File diffs: {result.file_diffs}\n"
            f"  Semantic diffs: {result.semantic_diffs}"
        )


class DescribeModificationPersistence:
    """Test that modifications persist through round-trip."""

    def it_preserves_added_paragraph(self, tmp_path: Path):
        """Test that adding a paragraph persists."""
        doc = Document()
        doc.add_paragraph("Original paragraph")

        path1 = tmp_path / "mod_para_v1.docx"
        doc.save(str(path1))

        # Reload and modify
        doc2 = Document(str(path1))
        marker = "ADDED_PARAGRAPH_MARKER_12345"
        doc2.add_paragraph(marker)

        path2 = tmp_path / "mod_para_v2.docx"
        doc2.save(str(path2))

        # Reload and check
        doc3 = Document(str(path2))
        texts = [p.text for p in doc3.paragraphs]

        assert marker in texts, f"Added paragraph not found. Found: {texts}"

    def it_preserves_modified_text(self, tmp_path: Path):
        """Test that modifying paragraph text persists."""
        doc = Document()
        doc.add_paragraph("Original text that will be changed")

        path1 = tmp_path / "mod_text_v1.docx"
        doc.save(str(path1))

        # Reload and modify
        doc2 = Document(str(path1))
        marker = "MODIFIED_TEXT_67890"
        doc2.paragraphs[0].text = marker

        path2 = tmp_path / "mod_text_v2.docx"
        doc2.save(str(path2))

        # Reload and check
        doc3 = Document(str(path2))

        assert doc3.paragraphs[0].text == marker

    def it_preserves_added_table(self, tmp_path: Path):
        """Test that adding a table persists."""
        doc = Document()
        doc.add_paragraph("Document with table")

        path1 = tmp_path / "mod_table_v1.docx"
        doc.save(str(path1))

        # Reload and add table
        doc2 = Document(str(path1))
        marker = "TABLE_CELL_MARKER_24680"
        table = doc2.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = marker

        path2 = tmp_path / "mod_table_v2.docx"
        doc2.save(str(path2))

        # Reload and check
        doc3 = Document(str(path2))

        assert len(doc3.tables) == 1
        assert doc3.tables[0].rows[0].cells[0].text == marker

    def it_preserves_formatting_changes(self, tmp_path: Path):
        """Test that formatting changes persist."""
        from docx.shared import Pt

        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("Text to format")

        path1 = tmp_path / "mod_format_v1.docx"
        doc.save(str(path1))

        # Reload and add formatting
        doc2 = Document(str(path1))
        run2 = doc2.paragraphs[0].runs[0]
        run2.bold = True
        run2.italic = True
        run2.font.size = Pt(24)

        path2 = tmp_path / "mod_format_v2.docx"
        doc2.save(str(path2))

        # Reload and check
        doc3 = Document(str(path2))
        run3 = doc3.paragraphs[0].runs[0]

        assert run3.bold is True
        assert run3.italic is True
        assert run3.font.size == Pt(24)


class DescribeEdgeCases:
    """Test edge cases and unusual combinations."""

    def it_handles_empty_document(self, tmp_path: Path):
        """Test that an empty document survives round-trip."""
        doc = Document()

        path1 = tmp_path / "empty_v1.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent

    def it_handles_single_empty_paragraph(self, tmp_path: Path):
        """Test that a document with a single empty paragraph survives."""
        doc = Document()
        doc.add_paragraph("")

        path1 = tmp_path / "single_empty_para.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent

    def it_handles_very_long_paragraph(self, tmp_path: Path):
        """Test that a very long paragraph survives."""
        doc = Document()
        long_text = "Lorem ipsum dolor sit amet. " * 500  # ~15KB of text
        doc.add_paragraph(long_text)

        path1 = tmp_path / "long_para.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent

        # Also verify content preserved
        doc2 = Document(str(tmp_path / "long_para.docx"))
        assert doc2.paragraphs[0].text == long_text

    def it_handles_many_tables(self, tmp_path: Path):
        """Test that many tables survive."""
        doc = Document()

        for i in range(20):
            doc.add_paragraph(f"Table {i+1}:")
            table = doc.add_table(rows=2, cols=3)
            for row in table.rows:
                for cell in row.cells:
                    cell.text = f"T{i+1}"

        path1 = tmp_path / "many_tables.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent

        # Verify count
        doc2 = Document(str(tmp_path / "many_tables.docx"))
        assert len(doc2.tables) == 20

    def it_handles_deep_table_nesting(self, tmp_path: Path):
        """Test that deeply nested tables survive."""
        doc = Document()

        # Create 3-level nested tables
        outer = doc.add_table(rows=1, cols=1)
        outer_cell = outer.rows[0].cells[0]
        outer_cell.paragraphs[0].text = "Outer"

        middle = outer_cell.add_table(rows=1, cols=1)
        middle_cell = middle.rows[0].cells[0]
        middle_cell.paragraphs[0].text = "Middle"

        inner = middle_cell.add_table(rows=1, cols=1)
        inner.rows[0].cells[0].text = "Inner"

        path1 = tmp_path / "deep_nested.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent

    def it_handles_special_characters_in_text(self, tmp_path: Path):
        """Test that special characters survive."""
        doc = Document()

        special = "Tab:\tNewline:\nBackslash:\\Quote:\"Ampersand:&Less:<Greater:>"
        doc.add_paragraph(special)

        path1 = tmp_path / "special_chars.docx"
        doc.save(str(path1))

        result = test_roundtrip_equivalence(path1, tmp_path)

        assert result.is_equivalent
